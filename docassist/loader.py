"""Чтение документов: PDF по страницам, Word и текст — целиком."""

from __future__ import annotations

import logging
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED = (".pdf", ".docx", ".txt", ".md")
# .docx — это zip: сжатый файл на 20 МБ может распаковаться в гигабайты и съесть память
MAX_DOCX_UNPACKED_MB = 200

log = logging.getLogger(__name__)


class UnreadableDocument(Exception):
    """Файл не читается или в нём нет текста (например, скан без распознавания)."""


@dataclass(frozen=True)
class Page:
    source: str  # имя файла, его видит пользователь в ссылке на источник
    page: int | None  # номер страницы PDF; у Word и текста страниц нет
    text: str


def read_document(path: Path) -> list[Page]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise UnreadableDocument(f"формат {suffix or 'без расширения'} не поддерживается")
    try:
        if suffix == ".pdf":
            pages = _read_pdf(path)
        elif suffix == ".docx":
            pages = [Page(path.name, None, _read_docx(path))]
        else:
            pages = [Page(path.name, None, path.read_text(encoding="utf-8", errors="replace"))]
    except UnreadableDocument:
        raise
    except Exception as e:  # битый файл любого формата — одна понятная ошибка
        # текст уходит в чат: подробности библиотеки (там бывает полный путь на сервере) — только в журнал
        log.warning("cannot read %s: %s", path.name, e)
        raise UnreadableDocument("файл не открылся — похоже, он повреждён") from e
    pages = [p for p in pages if p.text.strip()]
    if not pages:
        raise UnreadableDocument("в файле нет текста — похоже на скан")
    return pages


def _read_pdf(path: Path) -> list[Page]:
    reader = PdfReader(str(path))
    # строки PDF переносятся посреди фразы — склеиваем всю страницу в один абзац
    return [Page(path.name, i, _squash(page.extract_text() or "")) for i, page in enumerate(reader.pages, 1)]


def _read_docx(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        unpacked = sum(i.file_size for i in z.infolist())
    if unpacked > MAX_DOCX_UNPACKED_MB * 1024 * 1024:
        raise UnreadableDocument(f"после распаковки файл больше {MAX_DOCX_UNPACKED_MB} МБ")
    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(line for line in lines if line.strip())


def _squash(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()
