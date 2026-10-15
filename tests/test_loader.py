import pytest

from docassist.loader import UnreadableDocument, read_document
from tests.conftest import SAMPLES


def test_pdf_by_pages():
    pages = read_document(SAMPLES / "Договор на ремонт.pdf")
    assert [p.page for p in pages] == [1, 2, 3]
    assert "6 месяцев" in pages[2].text
    # переносы строк PDF склеены — фраза ищется целиком
    assert "на установленные новые детали" in pages[2].text


def test_docx_and_txt():
    (docx,) = read_document(SAMPLES / "Регламент возвратов.docx")
    assert docx.page is None and "service@technodom.example" in docx.text
    (txt,) = read_document(SAMPLES / "Памятка приёмщику.txt")
    assert "суббота с 10:00 до 18:00" in txt.text


def test_unreadable(tmp_path):
    empty = tmp_path / "скан.pdf"
    empty.write_bytes(b"%PDF-1.4 broken")
    with pytest.raises(UnreadableDocument):
        read_document(empty)
    blank = tmp_path / "пусто.txt"
    blank.write_text("  \n")
    with pytest.raises(UnreadableDocument, match="нет текста"):
        read_document(blank)
    with pytest.raises(UnreadableDocument, match="не поддерживается"):
        read_document(tmp_path / "таблица.xlsx")


def test_docx_tables_and_upper_suffix(tmp_path):
    from docx import Document

    path = tmp_path / "Прайс.DOCX"
    d = Document()
    d.add_paragraph("Цены на ремонт")
    d.add_paragraph("   ")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Услуга", "Цена"
    table.cell(1, 0).text, table.cell(1, 1).text = " Диагностика ", "500"
    d.save(path)
    (page,) = read_document(path)
    assert page.source == "Прайс.DOCX" and page.page is None
    assert page.text == "Цены на ремонт\nУслуга | Цена\nДиагностика | 500"


def test_markdown_and_bad_encoding(tmp_path):
    md = tmp_path / "faq.md"
    md.write_text("# Вопросы\nОтвет.", encoding="utf-8")
    assert read_document(md)[0].text == "# Вопросы\nОтвет."
    cp = tmp_path / "старый.txt"
    cp.write_bytes("Гарантия".encode("cp1251"))
    # чужая кодировка не роняет чтение
    assert read_document(cp)[0].text


def test_broken_docx(tmp_path):
    bad = tmp_path / "битый.docx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(UnreadableDocument, match="не открылся"):
        read_document(bad)
    with pytest.raises(UnreadableDocument, match="без расширения"):
        read_document(tmp_path / "README")
